from pathlib import Path
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter


repo_root = Path(__file__).resolve().parents[2]

results_dir = repo_root / "results"
results_dir.mkdir(parents=True, exist_ok=True)

output_excel = results_dir / "magnetic_materials_results_master_v2.xlsx"

paper_tables_dir = repo_root / "figures" / "paper" / "tables"
paper_tables_dir.mkdir(parents=True, exist_ok=True)

scenario_order = ["LOEO", "LOPO", "LOGO", "LOCO(k=10)"]
model_order = ["Random Forest", "XGBoost"]


def load_csv(path: Path, dataset_name=None):
    if not path.exists():
        print(f"[WARNING] Missing file: {path}")
        return pd.DataFrame()

    df = pd.read_csv(path)

    if dataset_name and "dataset" not in df.columns and "Dataset" not in df.columns:
        df.insert(0, "dataset", dataset_name)

    return df


def make_fig_mse_source(ood_t4: pd.DataFrame) -> pd.DataFrame:
    if ood_t4.empty:
        return pd.DataFrame()

    df = ood_t4.copy()
    df = df[
        df["scenario"].isin(scenario_order)
        & df["model"].isin(model_order)
    ].copy()

    df["MSE_mean"] = df["MSE"].astype(str).str.split("±").str[0].str.strip().astype(float)
    df["MSE_std"] = df["MSE"].astype(str).str.split("±").str[1].str.strip().astype(float)

    return df


def make_paper_table3(ood_t4: pd.DataFrame) -> pd.DataFrame:
    if ood_t4.empty:
        return pd.DataFrame()

    df = ood_t4.copy()
    df = df[
        df["scenario"].isin(scenario_order)
        & df["model"].isin(model_order)
    ].copy()

    df = df.rename(
        columns={
            "dataset": "Dataset",
            "scenario": "Scenario",
            "model": "Model",
        }
    )

    df["Dataset"] = df["Dataset"].replace(
        {
            "novamag": "Novamag",
            "mp": "Materials Project",
        }
    )

    return df[
        [
            "Dataset",
            "Scenario",
            "Model",
            "MSE",
            "MAE",
            "MRE",
            "R2",
            "n_splits",
        ]
    ]


def make_paper_table4(ood_t3: pd.DataFrame) -> pd.DataFrame:
    if ood_t3.empty:
        return pd.DataFrame()

    df = ood_t3.copy()
    df = df[df["scenario"].isin(scenario_order)].copy()

    summary = (
        df.groupby(["dataset", "scenario", "metric"], observed=True)
        .agg(
            **{
                "Significant / Total": (
                    "significant",
                    lambda x: f"{int(x.sum())} / {len(x)}",
                ),
                "Median t-test p": ("t_pvalue", "median"),
                "Median Wilcoxon p": ("wilcoxon_pvalue", "median"),
            }
        )
        .reset_index()
    )

    summary = summary.rename(
        columns={
            "dataset": "Dataset",
            "scenario": "Scenario",
            "metric": "Metric",
        }
    )

    summary["Dataset"] = summary["Dataset"].replace(
        {
            "novamag": "Novamag",
            "mp": "Materials Project",
        }
    )

    return summary


def make_id_vs_ood(cv_all: pd.DataFrame, ood_t4: pd.DataFrame) -> pd.DataFrame:
    if cv_all.empty or ood_t4.empty:
        return pd.DataFrame()

    cv = cv_all.copy()
    cv = cv[
        (cv["metric"] == "mse")
        & (cv["model"].isin(model_order))
    ].copy()

    cv = cv.rename(
        columns={
            "dataset": "Dataset",
            "model": "Model",
            "avg_runs_mean": "ID CV MSE mean",
            "avg_runs_std": "ID CV MSE std",
        }
    )

    cv["Dataset"] = cv["Dataset"].replace(
        {
            "novamag": "Novamag",
            "mp": "Materials Project",
        }
    )

    cv = cv[
        [
            "Dataset",
            "Model",
            "ID CV MSE mean",
            "ID CV MSE std",
        ]
    ]

    ood = ood_t4.copy()
    ood = ood[
        ood["scenario"].isin(scenario_order)
        & ood["model"].isin(model_order)
    ].copy()

    ood["MSE_mean"] = ood["MSE"].astype(str).str.split("±").str[0].str.strip().astype(float)

    ood_summary = (
        ood.groupby(["dataset", "model"], observed=True)
        .agg(
            **{
                "Mean OOD MSE": ("MSE_mean", "mean"),
                "OOD scenarios included": (
                    "scenario",
                    lambda x: ", ".join(sorted(set(x), key=scenario_order.index)),
                ),
            }
        )
        .reset_index()
        .rename(
            columns={
                "dataset": "Dataset",
                "model": "Model",
            }
        )
    )

    ood_summary["Dataset"] = ood_summary["Dataset"].replace(
        {
            "novamag": "Novamag",
            "mp": "Materials Project",
        }
    )

    merged = cv.merge(
        ood_summary,
        on=["Dataset", "Model"],
        how="left",
    )

    merged["Generalisation gap ΔMSE"] = (
        merged["Mean OOD MSE"] - merged["ID CV MSE mean"]
    )

    return merged


cv_novamag = load_csv(results_dir / "novamag_cv10_runs.csv", "novamag")
cv_mp = load_csv(results_dir / "mp_cv10_runs.csv", "mp")
cv_sig = load_csv(results_dir / "rf_vs_xgb_significance.csv")

ood_nova_t1 = load_csv(results_dir / "ood" / "table1_splits_summary.csv", "novamag")
ood_nova_t2 = load_csv(results_dir / "ood" / "table2_metrics_by_model.csv", "novamag")
ood_nova_t3 = load_csv(results_dir / "ood" / "table3_rf_vs_xgb_significance.csv", "novamag")
ood_nova_t4 = load_csv(results_dir / "ood" / "table4_combined_comparison.csv", "novamag")

ood_mp_t1 = load_csv(results_dir / "mp_ood" / "table1_splits_summary.csv", "mp")
ood_mp_t2 = load_csv(results_dir / "mp_ood" / "table2_metrics_by_model.csv", "mp")
ood_mp_t3 = load_csv(results_dir / "mp_ood" / "table3_rf_vs_xgb_significance.csv", "mp")
ood_mp_t4 = load_csv(results_dir / "mp_ood" / "table4_combined_comparison.csv", "mp")

cv_all = pd.concat([cv_novamag, cv_mp], ignore_index=True)

ood_t1 = pd.concat([ood_nova_t1, ood_mp_t1], ignore_index=True)
ood_t2 = pd.concat([ood_nova_t2, ood_mp_t2], ignore_index=True)
ood_t3 = pd.concat([ood_nova_t3, ood_mp_t3], ignore_index=True)
ood_t4 = pd.concat([ood_nova_t4, ood_mp_t4], ignore_index=True)

fig_mse = make_fig_mse_source(ood_t4)
paper_table3 = make_paper_table3(ood_t4)
paper_table4 = make_paper_table4(ood_t3)
id_vs_ood = make_id_vs_ood(cv_all, ood_t4)

readme = pd.DataFrame(
    {
        "Section": [
            "Purpose",
            "Main journal requirement",
            "ID CV",
            "OOD",
            "Figure source",
            "Paper Table 3",
            "Paper Table 4",
            "Traceability",
        ],
        "Description": [
            "Single master workbook containing all numerical results used in the paper.",
            "One clear, filterable, coloured source file for cross-validation, OOD results, figure source data, and paper-ready tables.",
            "In-distribution 10-fold cross-validation baseline. These results are not OOD.",
            "Structured held-out OOD splits: LOEO, LOPO, LOGO, and LOCO(k=10).",
            "Sheet 11 contains the exact MSE source data used for the two-panel OOD figure.",
            "Sheet 12A contains the combined OOD performance table used in the paper.",
            "Sheet 12B contains the combined RF vs XGB statistical significance table used in the paper.",
            "Each paper-facing table is generated directly from the raw CSV outputs, not manually edited.",
        ],
    }
)

notes = pd.DataFrame(
    {
        "Item": [
            "Dataset labels",
            "CV interpretation",
            "OOD interpretation",
            "MSE",
            "MAE",
            "MRE",
            "R2",
            "Significance",
            "Generalisation gap",
        ],
        "Note": [
            "Novamag and Materials Project are kept separate throughout the workbook.",
            "ID CV estimates interpolation performance under random splits within the same distribution.",
            "OOD estimates extrapolation performance under structured held-out chemical regions or clusters.",
            "Lower MSE is better.",
            "Lower MAE is better.",
            "Lower MRE is better.",
            "Higher R2 is better.",
            "RF vs XGB is summarised using paired t-test and Wilcoxon signed-rank test.",
            "Mean OOD MSE minus ID CV MSE mean.",
        ],
    }
)

with pd.ExcelWriter(output_excel, engine="openpyxl") as writer:
    readme.to_excel(writer, sheet_name="00_README", index=False)
    cv_novamag.to_excel(writer, sheet_name="01_ID_CV_Novamag", index=False)
    cv_mp.to_excel(writer, sheet_name="02_ID_CV_MP", index=False)

    ood_nova_t1.to_excel(writer, sheet_name="03_Nova_OOD_Splits", index=False)
    ood_nova_t2.to_excel(writer, sheet_name="04_Nova_OOD_Metrics", index=False)
    ood_nova_t3.to_excel(writer, sheet_name="05_Nova_OOD_Signif", index=False)
    ood_nova_t4.to_excel(writer, sheet_name="06_Nova_OOD_Combined", index=False)

    ood_mp_t1.to_excel(writer, sheet_name="07_MP_OOD_Splits", index=False)
    ood_mp_t2.to_excel(writer, sheet_name="08_MP_OOD_Metrics", index=False)
    ood_mp_t3.to_excel(writer, sheet_name="09_MP_OOD_Signif", index=False)
    ood_mp_t4.to_excel(writer, sheet_name="10_MP_OOD_Combined", index=False)

    fig_mse.to_excel(writer, sheet_name="11_Fig4_Source_MSE", index=False)
    paper_table3.to_excel(writer, sheet_name="12A_Paper_Table3_OOD", index=False)
    paper_table4.to_excel(writer, sheet_name="12B_Paper_Table4_Signif", index=False)
    notes.to_excel(writer, sheet_name="13_Notes", index=False)
    cv_sig.to_excel(writer, sheet_name="14_ID_CV_Significance", index=False)
    id_vs_ood.to_excel(writer, sheet_name="15_ID_vs_OOD", index=False)


wb = load_workbook(output_excel)

tab_colours = {
    "00_README": "1F4E78",
    "01_ID_CV_Novamag": "5B9BD5",
    "02_ID_CV_MP": "5B9BD5",
    "03_Nova_OOD_Splits": "70AD47",
    "04_Nova_OOD_Metrics": "70AD47",
    "05_Nova_OOD_Signif": "70AD47",
    "06_Nova_OOD_Combined": "70AD47",
    "07_MP_OOD_Splits": "F4B183",
    "08_MP_OOD_Metrics": "F4B183",
    "09_MP_OOD_Signif": "F4B183",
    "10_MP_OOD_Combined": "F4B183",
    "11_Fig4_Source_MSE": "9EADCC",
    "12A_Paper_Table3_OOD": "8064A2",
    "12B_Paper_Table4_Signif": "8064A2",
    "13_Notes": "A5A5A5",
    "14_ID_CV_Significance": "5B9BD5",
    "15_ID_vs_OOD": "C55A11",
}

header_fill = PatternFill("solid", fgColor="1F4E78")
paper_fill = PatternFill("solid", fgColor="8064A2")
fig_fill = PatternFill("solid", fgColor="9EADCC")
id_fill = PatternFill("solid", fgColor="5B9BD5")
ood_nova_fill = PatternFill("solid", fgColor="70AD47")
ood_mp_fill = PatternFill("solid", fgColor="F4B183")
light_fill = PatternFill("solid", fgColor="F7FBFF")

white_font = Font(color="FFFFFF", bold=True)
dark_font = Font(color="1F1F1F", bold=True)
body_font = Font(color="1F1F1F")
thin_border = Border(
    left=Side(style="thin", color="D9E2F3"),
    right=Side(style="thin", color="D9E2F3"),
    top=Side(style="thin", color="D9E2F3"),
    bottom=Side(style="thin", color="D9E2F3"),
)


def sheet_header_fill(sheet_name):
    if sheet_name.startswith("01") or sheet_name.startswith("02") or sheet_name.startswith("14"):
        return id_fill, white_font
    if "Nova_OOD" in sheet_name:
        return ood_nova_fill, white_font
    if "MP_OOD" in sheet_name:
        return ood_mp_fill, dark_font
    if sheet_name.startswith("11"):
        return fig_fill, dark_font
    if sheet_name.startswith("12"):
        return paper_fill, white_font
    if sheet_name.startswith("15"):
        return PatternFill("solid", fgColor="C55A11"), white_font
    return header_fill, white_font


def format_worksheet(ws):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    max_row = ws.max_row
    max_col = ws.max_column

    if max_row < 1 or max_col < 1:
        return

    fill, font = sheet_header_fill(ws.title)

    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.font = body_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if cell.row % 2 == 0:
                cell.fill = light_fill

    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx)
        values = [
            str(ws.cell(row=r, column=col_idx).value)
            for r in range(1, min(max_row, 80) + 1)
            if ws.cell(row=r, column=col_idx).value is not None
        ]
        width = min(max(max((len(v) for v in values), default=10) + 2, 10), 38)
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 24

    table_ref = f"A1:{get_column_letter(max_col)}{max_row}"
    clean_name = (
        "T_" + ws.title.replace(" ", "_").replace("-", "_").replace("(", "").replace(")", "")
    )[:250]

    if max_row >= 2 and max_col >= 1:


     ws.auto_filter.ref = table_ref

    for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"


def add_conditional_formatting(ws):
    if ws.max_row < 3:
        return

    headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}

    for metric_col in [
        "MSE_mean",
        "MSE_std",
        "MSE mean",
        "MSE std",
        "Mean OOD MSE",
        "ID CV MSE mean",
        "Generalisation gap ΔMSE",
        "Median t-test p",
        "Median Wilcoxon p",
    ]:
        if metric_col in headers:
            col = get_column_letter(headers[metric_col])
            rng = f"{col}2:{col}{ws.max_row}"
            ws.conditional_formatting.add(
                rng,
                ColorScaleRule(
                    start_type="min",
                    start_color="63BE7B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="F8696B",
                ),
            )

    for metric_col in ["R2", "r2", "avg_runs_mean"]:
        if metric_col in headers:
            col = get_column_letter(headers[metric_col])
            rng = f"{col}2:{col}{ws.max_row}"
            ws.conditional_formatting.add(
                rng,
                ColorScaleRule(
                    start_type="min",
                    start_color="F8696B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="63BE7B",
                ),
            )


for ws in wb.worksheets:
    ws.sheet_properties.tabColor = tab_colours.get(ws.title, "A5A5A5")
    format_worksheet(ws)
    add_conditional_formatting(ws)

wb.save(output_excel)

paper_table3.to_csv(
    paper_tables_dir / "table3_combined_ood_performance.csv",
    index=False,
)
paper_table4.to_csv(
    paper_tables_dir / "table4_combined_rf_xgb_significance.csv",
    index=False,
)
fig_mse.to_csv(
    paper_tables_dir / "figure4_source_mse.csv",
    index=False,
)
id_vs_ood.to_csv(
    paper_tables_dir / "id_vs_ood_comparison.csv",
    index=False,
)

print(f"\nMaster Excel file created at:\n{output_excel}")
print(f"Paper Table 3 CSV:\n{paper_tables_dir / 'table3_combined_ood_performance.csv'}")
print(f"Paper Table 4 CSV:\n{paper_tables_dir / 'table4_combined_rf_xgb_significance.csv'}")
print(f"Figure 4 source CSV:\n{paper_tables_dir / 'figure4_source_mse.csv'}")
print(f"ID vs OOD comparison CSV:\n{paper_tables_dir / 'id_vs_ood_comparison.csv'}")